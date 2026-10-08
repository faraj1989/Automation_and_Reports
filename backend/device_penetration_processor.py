#!/usr/bin/env python3
"""
Libyana NPM - Weekly Device Penetration Reader / Archive / Weekly Report

Reads Weekly_Device_Penetration_Historical.*, produced by
scrapers/weekly_device_penetration_scraper.py (same SmartCare portal/login
as the CEM pipeline in smartcare_cem_processor.py, different dashboard).

One row per device model per daily snapshot ('Time'): Time, Device Model,
Device Brand, Device Type, Device OS, Device Technology, Number of
Users(count), Penetration Rate(%). Network-wide, no per-site breakdown.
Each weekly export is a rolling 7 days (~45k rows), so a missed run is a
permanent gap - the portal can't be asked for older days.

Archive: the Parquet file is the full history (no row cap, ~1s to read).
The .xlsx alongside it is only a rolling mirror of the last
XLSX_MIRROR_WEEKS weeks - a full-history workbook would hit Excel's
1,048,576-row sheet limit after ~6 months and stop updating.
"""
import io
import os
import logging
from datetime import datetime, timedelta
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import pandas as pd

logger = logging.getLogger(__name__)

DEVICE_PENETRATION_PATH = os.environ.get(
    "DEVICE_PENETRATION_HISTORY_PATH",
    r"C:\Users\user\Desktop\Libyana_Data\Output\Weekly_Device_Penetration_Exports\Weekly_Device_Penetration_Historical.xlsx",
)
HISTORICAL_BASENAME = "Weekly_Device_Penetration_Historical"
REPORTS_SUBDIR = "Reports"
XLSX_MIRROR_WEEKS = 8

# Same weekly cadence as SmartCare CEM (schedule_config.json: Sunday 00:00)
# - see smartcare_cem_processor.CEM_STALE_DAYS for the same reasoning.
DEVICE_PENETRATION_STALE_DAYS = 10

USERS = 'Number of Users(count)'
IDENTITY_COLUMNS = ["Device Model", "Device Brand", "Device Type", "Device OS", "Device Technology"]
HISTORY_DEDUPE_COLUMNS = ["Time"] + IDENTITY_COLUMNS
HISTORY_COLUMNS = HISTORY_DEDUPE_COLUMNS + [USERS, "Penetration Rate(%)"]

# IoT/POS modules (one QUECTEL module is ~6% of all devices) and routers
# swamp brand/model rankings, so those default to subscriber handsets only.
HANDSET_TYPES = ('SmartPhone', 'FeaturePhone', 'Tablet')

# 4G/5G-capable vs. legacy-only, derived from the 'Device Technology'
# column (e.g. "2G/3G/LTE", "2G/3G/LTE/NR", "2G", "3G/LTE") - used for a
# quick network-wide "how much of the device base could actually use a
# newly activated LTE band" read, complementing the per-site 4G-band
# capacity advice in ReportGenerator.build_site_capacity_advice().
BROADBAND_TERMS = ("LTE", "NR")
TECH_ORDER = ['5G NR', 'LTE (no 5G)', '3G (no LTE)', '2G only', 'Other', 'Unidentified']


def tech_bucket(tech: pd.Series) -> pd.Series:
    """Bucket 'Device Technology' strings by the newest RAT the device
    supports. Vectorized - this runs over the full multi-week history."""
    t = tech.fillna('').astype(str)
    bucket = pd.Series('Other', index=t.index)
    bucket[t.str.contains('2G')] = '2G only'
    bucket[t.str.contains('3G')] = '3G (no LTE)'
    bucket[t.str.contains('LTE')] = 'LTE (no 5G)'
    bucket[t.str.contains('NR')] = '5G NR'
    bucket[t.str.lower().isin(['', 'unidentified', 'unknown'])] = 'Unidentified'
    return bucket


# ---------------------------- Archive ----------------------------

def _paths(output_dir=None):
    base = Path(output_dir) if output_dir else Path(DEVICE_PENETRATION_PATH).parent
    return base / f"{HISTORICAL_BASENAME}.parquet", base / f"{HISTORICAL_BASENAME}.xlsx"


def load_history(output_dir=None) -> Optional[pd.DataFrame]:
    """Full history - Parquet if present, else the legacy full .xlsx (the
    first fold after this change migrates it to Parquet)."""
    parquet_path, xlsx_path = _paths(output_dir)
    if parquet_path.exists():
        return pd.read_parquet(parquet_path)
    if xlsx_path.exists():
        return pd.read_excel(xlsx_path, engine='openpyxl')
    return None


def write_history(df: pd.DataFrame, output_dir=None) -> Path:
    """Atomic Parquet write (full history) + a rolling .xlsx mirror of the
    last XLSX_MIRROR_WEEKS weeks for people who open it by hand."""
    parquet_path, xlsx_path = _paths(output_dir)
    df = df.sort_values(['Time', USERS], ascending=[True, False]).reset_index(drop=True)
    df['Time'] = df['Time'].astype(str)

    tmp = parquet_path.with_suffix('.parquet.tmp')
    df.to_parquet(tmp, index=False)
    os.replace(tmp, parquet_path)

    cutoff = pd.to_datetime(df['Time']).max().normalize() - timedelta(weeks=XLSX_MIRROR_WEEKS)
    mirror = df[pd.to_datetime(df['Time']) > cutoff]
    tmp = xlsx_path.with_name(xlsx_path.stem + '.tmp.xlsx')
    mirror.to_excel(tmp, index=False)
    os.replace(tmp, xlsx_path)
    return parquet_path


def fold_export(new_data: pd.DataFrame, output_dir=None) -> pd.DataFrame:
    """Merge one export into the archive, deduped by snapshot Time + device
    identity (keep='last') so re-processing an export never double-counts.
    Returns the new rows (for picking which weekly reports to regenerate)."""
    missing = [c for c in HISTORY_COLUMNS if c not in new_data.columns]
    if missing:
        raise ValueError(f"not a Device Penetration export (missing columns: {missing}) - history left unchanged")
    new_data = new_data.dropna(subset=["Device Model"])[HISTORY_COLUMNS].copy()
    new_data['Time'] = new_data['Time'].astype(str)
    history = load_history(output_dir)
    combined = new_data if history is None else pd.concat([history, new_data], ignore_index=True)
    combined = combined.drop_duplicates(subset=HISTORY_DEDUPE_COLUMNS, keep="last")
    path = write_history(combined, output_dir)
    logger.info(f"Historical dataset updated: {path} ({len(combined)} total rows, "
                f"{combined['Time'].nunique()} daily snapshots)")
    return new_data


# ---------------------------- Dashboard overview ----------------------------

def is_available() -> bool:
    return any(p.exists() for p in _paths())


def _archive_mtime() -> Optional[float]:
    existing = [p for p in _paths() if p.exists()]
    return max(os.path.getmtime(p) for p in existing) if existing else None


def load_device_penetration_overview() -> Dict:
    """Full history plus the latest snapshot's breakdowns by Device Type,
    Device Brand, and broadband (LTE/NR) capability - the scraper reruns
    weekly, so a stale file here means a missed run."""
    try:
        df = load_history()
    except Exception as e:
        logger.warning(f"Could not read Device Penetration history: {e}")
        return {'loaded': False}
    if df is None or df.empty or 'Time' not in df.columns:
        return {'loaded': False}
    df['Time'] = df['Time'].astype(str)

    latest_time = df['Time'].max()
    latest = df[df['Time'] == latest_time].copy()
    if 'Device Technology' in latest.columns:
        latest['Broadband Capable'] = latest['Device Technology'].astype(str).str.contains(
            '|'.join(BROADBAND_TERMS), na=False)

    mtime = _archive_mtime()
    age_days = (datetime.now().timestamp() - mtime) / 86400.0
    return {
        'loaded': True,
        'age_days': age_days,
        'is_stale': age_days > DEVICE_PENETRATION_STALE_DAYS,
        'last_updated': datetime.fromtimestamp(mtime),
        'latest_time': latest_time,
        'latest_snapshot': latest,
        'history': df,
    }


# ---------------------------- Weekly report ----------------------------

def _week_start(times: pd.Series) -> pd.Series:
    d = pd.to_datetime(times).dt.normalize()
    return d - pd.to_timedelta(d.dt.weekday, unit='D')  # Monday


def week_label(week_start, days: Optional[int] = None) -> str:
    ws = pd.Timestamp(week_start)
    iso = ws.isocalendar()
    label = f"{iso.year}-W{iso.week:02d} ({ws:%b %d} – {ws + timedelta(days=6):%b %d})"
    if days is not None and days < 7:
        label += f" · {days}/7 days"
    return label


def available_weeks(history: pd.DataFrame) -> pd.DataFrame:
    """One row per Monday-Sunday week in the archive, newest first."""
    days = pd.to_datetime(history['Time']).dt.normalize().drop_duplicates()
    weeks = pd.DataFrame({'day': days, 'week_start': _week_start(days)})
    out = weeks.groupby('week_start')['day'].nunique().rename('days').reset_index()
    out['label'] = [week_label(w, d) for w, d in zip(out['week_start'], out['days'])]
    return out.sort_values('week_start', ascending=False).reset_index(drop=True)


def _week_slice(history: pd.DataFrame, week_start) -> Tuple[pd.DataFrame, int]:
    ws = pd.Timestamp(week_start)
    t = pd.to_datetime(history['Time'])
    rows = history[(t >= ws) & (t < ws + timedelta(days=7))]
    return rows, rows['Time'].nunique()


def _avg_daily(rows: pd.DataFrame, days: int, by) -> pd.Series:
    # Users are daily unique counts - not additive across days - so the
    # weekly figure is the average day. Divide by days (not .mean()) so a
    # model absent on some days counts as 0 for those days.
    return rows.groupby(by)[USERS].sum() / max(days, 1)


def _kpis(rows: pd.DataFrame, days: int) -> Dict:
    total = rows[USERS].sum() / max(days, 1)
    tech = _avg_daily(rows, days, tech_bucket(rows['Device Technology']))
    pct = lambda *buckets: 100 * sum(tech.get(b, 0) for b in buckets) / total if total else float('nan')
    handset = rows[rows['Device Type'].isin(HANDSET_TYPES)][USERS].sum() / max(days, 1)
    return {
        'Avg Daily Devices': round(total),
        'LTE/NR-Capable %': pct('5G NR', 'LTE (no 5G)'),
        '5G (NR)-Capable %': pct('5G NR'),
        '2G-Only %': pct('2G only'),
        'Handset Share %': 100 * handset / total if total else float('nan'),
        'Distinct Models (week)': rows['Device Model'].nunique(),
    }


def _share_table(cur: pd.Series, prev: Optional[pd.Series], name: str, top: Optional[int] = None) -> pd.DataFrame:
    cur = cur.sort_values(ascending=False)
    df = pd.DataFrame({name: cur.index, 'Avg Daily Users': cur.values.round(0),
                       'Share %': (100 * cur / cur.sum()).values.round(2)})
    if prev is not None and not prev.empty:
        prev_share = (100 * prev / prev.sum()).round(2)
        df['Prev Share %'] = df[name].map(prev_share)
        df['Δ Share (pp)'] = (df['Share %'] - df['Prev Share %']).round(2)
    if top:
        df = df.head(top)
    df.insert(0, 'Rank', range(1, len(df) + 1))
    return df


def weekly_summary(history: pd.DataFrame, week_start, handsets_only: bool = True) -> Dict:
    """Everything the weekly report shows - KPIs vs previous week, ranked
    brand/model/OS tables, all-device type/RAT mix, and biggest movers.
    'Previous week' is the newest earlier week in the archive, flagged
    when it isn't the adjacent one (e.g. across the Sep 10-27 gap)."""
    history = history.copy()
    history['Time'] = history['Time'].astype(str)
    weeks = available_weeks(history)
    ws = pd.Timestamp(week_start)
    cur, days = _week_slice(history, ws)
    earlier = weeks[weeks['week_start'] < ws]
    prev_ws = earlier['week_start'].iloc[0] if not earlier.empty else None
    prev, prev_days = _week_slice(history, prev_ws) if prev_ws is not None else (cur.iloc[0:0], 0)

    def ranked(rows):
        return rows[rows['Device Type'].isin(HANDSET_TYPES)] if handsets_only else rows

    def by(rows, n_days, col):
        return _avg_daily(rows, n_days, col) if n_days else None

    rc, rp = ranked(cur), ranked(prev)
    model_cur = by(rc, days, 'Device Model')
    model_prev = by(rp, prev_days, 'Device Model')
    models = _share_table(model_cur, model_prev, 'Device Model')
    attrs = (rc.sort_values(USERS, ascending=False)
             .drop_duplicates('Device Model').set_index('Device Model')[IDENTITY_COLUMNS[1:]])
    models = models.join(attrs, on='Device Model')
    models = models[['Rank', 'Device Model'] + IDENTITY_COLUMNS[1:] +
                    [c for c in models.columns if c not in ['Rank', 'Device Model'] + IDENTITY_COLUMNS[1:]]]

    movers = pd.DataFrame()
    if model_prev is not None:
        delta = model_cur.sub(model_prev, fill_value=0).round(0)
        movers = pd.DataFrame({'Device Model': delta.index, 'Δ Avg Daily Users': delta.values,
                               'This Week': model_cur.reindex(delta.index).fillna(0).round(0).values,
                               'Prev Week': model_prev.reindex(delta.index).fillna(0).round(0).values})
        movers = movers.join(attrs[['Device Brand']], on='Device Model')
        gainers = movers.nlargest(15, 'Δ Avg Daily Users').assign(Movement='▲ Gainer')
        losers = movers.nsmallest(15, 'Δ Avg Daily Users').assign(Movement='▼ Loser')
        movers = pd.concat([gainers, losers]).drop_duplicates('Device Model')
        movers = movers[['Movement'] + [c for c in movers.columns if c != 'Movement']]

    tech_cur = _avg_daily(cur, days, tech_bucket(cur['Device Technology'])).reindex(TECH_ORDER).dropna()
    tech_prev = (_avg_daily(prev, prev_days, tech_bucket(prev['Device Technology'])) if prev_days else None)
    tech = _share_table(tech_cur, tech_prev, 'Highest Supported RAT')
    tech = tech.set_index('Highest Supported RAT').reindex([b for b in TECH_ORDER if b in tech_cur.index]).reset_index()
    tech['Rank'] = range(1, len(tech) + 1)

    return {
        'week_start': ws, 'days': days, 'label': week_label(ws, days),
        'prev_week_start': prev_ws, 'prev_days': prev_days,
        'prev_label': week_label(prev_ws, prev_days) if prev_ws is not None else None,
        'prev_is_adjacent': prev_ws is not None and (ws - prev_ws).days == 7,
        'handsets_only': handsets_only,
        'kpis': _kpis(cur, days),
        'prev_kpis': _kpis(prev, prev_days) if prev_days else None,
        'brands': _share_table(by(rc, days, 'Device Brand'), by(rp, prev_days, 'Device Brand'), 'Device Brand'),
        'models': models,
        'os': _share_table(by(rc, days, 'Device OS'), by(rp, prev_days, 'Device OS'), 'Device OS'),
        'types': _share_table(by(cur, days, 'Device Type'), by(prev, prev_days, 'Device Type'), 'Device Type'),
        'tech': tech,
        'movers': movers,
    }


def build_weekly_report_xlsx(summary: Dict) -> bytes:
    """Formatted weekly report workbook with native Excel charts."""
    from openpyxl.chart import BarChart, PieChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill('solid', fgColor='1F4E78')
    header_font = Font(bold=True, color='FFFFFF')
    scope = "Handsets only (SmartPhone / FeaturePhone / Tablet)" if summary['handsets_only'] else "All devices"

    kpis, prev = summary['kpis'], summary['prev_kpis']
    kpi_rows = []
    for k, v in kpis.items():
        p = prev.get(k) if prev else None
        kpi_rows.append({'KPI': k, 'This Week': round(v, 2),
                         'Previous Week': round(p, 2) if p is not None else None,
                         'Change': round(v - p, 2) if p is not None else None})

    sheets = [
        ('Summary', pd.DataFrame(kpi_rows)),
        ('Top Brands', summary['brands'].head(30)),
        ('Top Models', summary['models'].head(50)),
        ('Movers', summary['movers']),
        ('Device OS', summary['os']),
        ('Device Type', summary['types']),
        ('Technology', summary['tech']),
        ('All Models', summary['models']),
    ]
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as writer:
        start = {'Summary': 7}
        for name, df in sheets:
            if df is None or df.empty:
                df = pd.DataFrame([{'Status': 'No data (no comparable previous week)'}])
            df.to_excel(writer, sheet_name=name, index=False, startrow=start.get(name, 0))
        wb = writer.book

        ws = wb['Summary']
        ws['A1'] = 'Libyana — Weekly Device Penetration Report'
        ws['A1'].font = Font(bold=True, size=14)
        ws['A2'] = f"Week: {summary['label']}"
        ws['A3'] = (f"Compared with: {summary['prev_label']}" +
                    ("" if summary['prev_is_adjacent'] else "  (not the adjacent week - archive gap)")
                    if summary['prev_label'] else "Compared with: no earlier week in archive")
        ws['A4'] = f"Brand / Model / OS scope: {scope}. Device Type and Technology: all devices."
        ws['A5'] = ("Users are daily unique counts; weekly values are the average day. "
                    f"Generated {datetime.now():%Y-%m-%d %H:%M}.")
        for r in range(2, 6):
            ws[f'A{r}'].font = Font(italic=r > 2, bold=r == 2, color='444444')

        for name, df in sheets:
            ws = wb[name]
            hdr = start.get(name, 0) + 1
            for cell in ws[hdr]:
                cell.fill, cell.font = header_fill, header_font
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            ws.freeze_panes = ws.cell(row=hdr + 1, column=1)
            for col in ws.iter_cols(min_row=hdr):
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width + 2, 10), 45)
                for c in col[1:]:
                    if isinstance(c.value, float):
                        c.number_format = '#,##0.00'
                    elif isinstance(c.value, int):
                        c.number_format = '#,##0'

        def bar(sheet, cat_col, val_col, n, title, anchor, horizontal=True):
            ws = wb[sheet]
            n = min(n, ws.max_row - 1)
            if n < 1:
                return
            ch = BarChart()
            ch.type = 'bar' if horizontal else 'col'
            ch.title, ch.legend, ch.height, ch.width = title, None, 9 if not horizontal else 0.55 * n + 3, 18
            ch.add_data(Reference(ws, min_col=val_col, min_row=1, max_row=n + 1), titles_from_data=True)
            ch.set_categories(Reference(ws, min_col=cat_col, min_row=2, max_row=n + 1))
            if horizontal:
                ch.y_axis.scaling.orientation = 'minMax'
                ch.x_axis.scaling.orientation = 'maxMin'  # rank 1 at top
            ch.dataLabels = DataLabelList(showVal=True)
            ws.add_chart(ch, anchor)

        cols = lambda sheet: [c.value for c in wb[sheet][1]]
        if 'Share %' in cols('Top Brands'):
            bar('Top Brands', 2, cols('Top Brands').index('Share %') + 1, 10, 'Top 10 Brands — share %',
                f"{get_column_letter(len(cols('Top Brands')) + 2)}2")
        if 'Share %' in cols('Top Models'):
            bar('Top Models', 2, cols('Top Models').index('Share %') + 1, 15, 'Top 15 Models — share %',
                f"{get_column_letter(len(cols('Top Models')) + 2)}2")
        if 'Share %' in cols('Technology'):
            bar('Technology', 2, cols('Technology').index('Share %') + 1, 6, 'Highest supported RAT — share %',
                f"{get_column_letter(len(cols('Technology')) + 2)}2", horizontal=False)
        os_ws = wb['Device OS']
        if 'Avg Daily Users' in cols('Device OS') and os_ws.max_row > 1:
            pie = PieChart()
            pie.title, pie.height, pie.width = 'Device OS', 9, 12
            n = min(5, os_ws.max_row - 1)
            val = cols('Device OS').index('Avg Daily Users') + 1
            pie.add_data(Reference(os_ws, min_col=val, min_row=1, max_row=n + 1), titles_from_data=True)
            pie.set_categories(Reference(os_ws, min_col=2, min_row=2, max_row=n + 1))
            pie.dataLabels = DataLabelList(showPercent=True)
            os_ws.add_chart(pie, f"{get_column_letter(len(cols('Device OS')) + 2)}2")
    buf.seek(0)
    return buf.read()


def report_filename(week_start) -> str:
    iso = pd.Timestamp(week_start).isocalendar()
    return f"Device_Penetration_Weekly_Report_{iso.year}-W{iso.week:02d}.xlsx"


def save_weekly_reports(history: pd.DataFrame, week_starts, output_dir=None) -> List[Path]:
    """Write/overwrite the report for each given week (a week touched by a
    new export is regenerated, so a partial week fills in on the next run)."""
    reports_dir = (Path(output_dir) if output_dir else Path(DEVICE_PENETRATION_PATH).parent) / REPORTS_SUBDIR
    reports_dir.mkdir(parents=True, exist_ok=True)
    written = []
    for ws in sorted(set(pd.Timestamp(w) for w in week_starts)):
        path = reports_dir / report_filename(ws)
        path.write_bytes(build_weekly_report_xlsx(weekly_summary(history, ws)))
        written.append(path)
        logger.info(f"Weekly report written: {path}")
    return written


def weeks_in(rows: pd.DataFrame) -> List[pd.Timestamp]:
    return list(_week_start(rows['Time'].drop_duplicates()).unique())


# ---------------------------- CLI ----------------------------
if __name__ == "__main__":
    import argparse
    logging.basicConfig(level=logging.INFO)
    parser = argparse.ArgumentParser(description="Device Penetration archive / weekly reports")
    parser.add_argument("--migrate", action="store_true", help="convert the legacy full .xlsx history to Parquet")
    parser.add_argument("--reports", action="store_true", help="(re)generate the weekly report for every archived week")
    args = parser.parse_args()

    if args.migrate:
        hist = load_history()
        if hist is not None:
            write_history(hist)
            print(f"Migrated {len(hist):,} rows to {_paths()[0]}")
    if args.reports:
        hist = load_history()
        for p in save_weekly_reports(hist, available_weeks(hist)['week_start']):
            print("Wrote", p)
    if not (args.migrate or args.reports):
        overview = load_device_penetration_overview()
        print("Loaded:", overview.get('loaded'), "| age (days):", overview.get('age_days'))
        if overview.get('loaded'):
            print(available_weeks(overview['history']).to_string(index=False))
